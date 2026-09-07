from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(slots=True)
class FileInspection:
    path: Path
    headers: list[str]
    has_data: bool
    sheet: str | None = None
    delimiter: str | None = None
    error: str | None = None


def inspect_file(path: Path, *, encoding: str = "utf-8-sig", delimiter: str | None = None) -> FileInspection:
    path = Path(path)
    try:
        if path.suffix.lower() == ".csv":
            return _inspect_csv(path, encoding=encoding, delimiter=delimiter)
        if path.suffix.lower() == ".xlsx":
            return _inspect_xlsx(path)
        return FileInspection(path=path, headers=[], has_data=False, error="Unsupported file type")
    except Exception as exc:
        return FileInspection(path=path, headers=[], has_data=False, error=f"{type(exc).__name__}: {exc}")


def _detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _inspect_csv(path: Path, *, encoding: str, delimiter: str | None) -> FileInspection:
    with path.open("r", encoding=encoding, newline="") as handle:
        sample = handle.read(64 * 1024)
        handle.seek(0)
        actual_delimiter = delimiter or _detect_delimiter(sample)
        reader = csv.reader(handle, delimiter=actual_delimiter)
        try:
            header = next(reader)
        except StopIteration:
            return FileInspection(path=path, headers=[], has_data=False, delimiter=actual_delimiter)
        headers = ["" if value is None else str(value) for value in header]
        has_data = any(any(str(cell).strip() for cell in row) for _, row in zip(range(10), reader))
        return FileInspection(path=path, headers=headers, has_data=has_data, delimiter=actual_delimiter)


def _inspect_xlsx(path: Path) -> FileInspection:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return FileInspection(path=path, headers=[], has_data=False, sheet=ws.title)
        headers = ["" if value is None else str(value) for value in header]
        has_data = False
        for _, row in zip(range(10), rows):
            if any(value is not None and str(value).strip() != "" for value in row):
                has_data = True
                break
        return FileInspection(path=path, headers=headers, has_data=has_data, sheet=ws.title)
    finally:
        wb.close()
