from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import tempfile
import shutil

import duckdb
from openpyxl import load_workbook

from data_validator.core.config import ValidationConfig
from data_validator.core.normalisation import normalise_qr, normalise_serial
from data_validator.ingestion.inspection import FileInspection
from data_validator.system.resources import ResourcePlan


def _sql_string(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _sql_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


@dataclass(slots=True)
class IngestSummary:
    file_id: int
    path: Path
    rows: int
    sheet: str | None = None


class ValidationDatabase:
    def __init__(self, plan: ResourcePlan, config: ValidationConfig):
        self.plan = plan
        self.config = config
        self.root = Path(tempfile.mkdtemp(prefix="data_validator_", dir=plan.temp_directory))
        self.db_path = self.root / "validation.duckdb"
        self.conn = duckdb.connect(str(self.db_path))
        self._configure()
        self._create_schema()

    def _configure(self) -> None:
        self.conn.execute(f"SET memory_limit = '{int(self.plan.memory_limit_mb)}MB'")
        self.conn.execute(f"SET threads = {int(self.plan.threads)}")
        spill = self.root / "spill"
        spill.mkdir(parents=True, exist_ok=True)
        self.conn.execute(f"SET temp_directory = {_sql_string(str(spill))}")

    def _create_schema(self) -> None:
        self.conn.execute(
            """
            CREATE TABLE records (
                file_id INTEGER NOT NULL,
                source_file VARCHAR NOT NULL,
                source_sheet VARCHAR,
                source_row BIGINT NOT NULL,
                sku_raw VARCHAR,
                serial_raw VARCHAR,
                qr_raw VARCHAR,
                sku VARCHAR NOT NULL,
                serial VARCHAR NOT NULL,
                qr VARCHAR NOT NULL,
                serial_num BIGINT
            )
            """
        )

    def close(self, cleanup: bool = True) -> None:
        try:
            self.conn.close()
        finally:
            if cleanup:
                shutil.rmtree(self.root, ignore_errors=True)

    def interrupt(self) -> None:
        try:
            self.conn.interrupt()
        except Exception:
            pass

    def ingest(self, file_id: int, inspection: FileInspection) -> IngestSummary:
        if inspection.path.suffix.lower() == ".csv":
            return self._ingest_csv(file_id, inspection)
        if inspection.path.suffix.lower() == ".xlsx":
            return self._ingest_xlsx(file_id, inspection)
        raise ValueError(f"Unsupported input type: {inspection.path}")

    def _normalised_sql(self, identifier: str, *, kind: str) -> str:
        raw = f"COALESCE(CAST({_sql_identifier(identifier)} AS VARCHAR), '')"
        if kind == "sku":
            return f"trim({raw})"
        if kind == "serial":
            expr = f"trim({raw})" if self.config.serial_trim_whitespace else raw
            if not self.config.serial_case_sensitive:
                expr = f"upper({expr})"
            return expr
        if kind == "qr":
            expr = f"trim({raw})" if self.config.qr_trim_whitespace else raw
            if not self.config.qr_case_sensitive:
                expr = f"upper({expr})"
            return expr
        raise ValueError(kind)

    def _ingest_csv(self, file_id: int, inspection: FileInspection) -> IngestSummary:
        c = self.config.columns
        delimiter = inspection.delimiter or self.config.csv_delimiter or ","
        source = _sql_string(str(inspection.path))
        delim = _sql_string(delimiter)
        sku_id = _sql_identifier(c.sku)
        serial_id = _sql_identifier(c.serial)
        qr_id = _sql_identifier(c.qr)
        sku_norm = self._normalised_sql(c.sku, kind="sku")
        serial_norm = self._normalised_sql(c.serial, kind="serial")
        qr_norm = self._normalised_sql(c.qr, kind="qr")
        sql = f"""
            INSERT INTO records
            SELECT
                {int(file_id)} AS file_id,
                {_sql_string(str(inspection.path))} AS source_file,
                NULL AS source_sheet,
                row_number() OVER () + 1 AS source_row,
                CAST({sku_id} AS VARCHAR) AS sku_raw,
                CAST({serial_id} AS VARCHAR) AS serial_raw,
                CAST({qr_id} AS VARCHAR) AS qr_raw,
                {sku_norm} AS sku,
                {serial_norm} AS serial,
                {qr_norm} AS qr,
                NULL::BIGINT AS serial_num
            FROM read_csv({source}, header=true, all_varchar=true, delim={delim}, null_padding=true)
        """
        self.conn.execute(sql)
        rows = self.conn.execute("SELECT count(*) FROM records WHERE file_id = ?", [file_id]).fetchone()[0]
        return IngestSummary(file_id=file_id, path=inspection.path, rows=int(rows))

    def _ingest_xlsx(self, file_id: int, inspection: FileInspection) -> IngestSummary:
        wb = load_workbook(inspection.path, read_only=True, data_only=True)
        try:
            sheet_name = self.config.xlsx_sheet or inspection.sheet or wb.sheetnames[0]
            ws = wb[sheet_name]
            iterator = ws.iter_rows(values_only=True)
            header = next(iterator)
            header_map = {str(value): idx for idx, value in enumerate(header) if value is not None}
            c = self.config.columns
            idx_sku = header_map[c.sku]
            idx_serial = header_map[c.serial]
            idx_qr = header_map[c.qr]
            batch: list[tuple] = []
            rows = 0
            insert_sql = """
                INSERT INTO records
                (file_id, source_file, source_sheet, source_row, sku_raw, serial_raw, qr_raw, sku, serial, qr, serial_num)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL)
            """
            for excel_row, values in enumerate(iterator, start=2):
                if not any(value is not None and str(value) != "" for value in values):
                    continue
                sku_raw = "" if idx_sku >= len(values) or values[idx_sku] is None else str(values[idx_sku])
                serial_raw = "" if idx_serial >= len(values) or values[idx_serial] is None else str(values[idx_serial])
                qr_raw = "" if idx_qr >= len(values) or values[idx_qr] is None else str(values[idx_qr])
                batch.append((
                    file_id,
                    str(inspection.path),
                    sheet_name,
                    excel_row,
                    sku_raw,
                    serial_raw,
                    qr_raw,
                    sku_raw.strip(),
                    normalise_serial(serial_raw, self.config),
                    normalise_qr(qr_raw, self.config),
                ))
                rows += 1
                if len(batch) >= self.config.xlsx_batch_size:
                    self.conn.executemany(insert_sql, batch)
                    batch.clear()
            if batch:
                self.conn.executemany(insert_sql, batch)
            return IngestSummary(file_id=file_id, path=inspection.path, rows=rows, sheet=sheet_name)
        finally:
            wb.close()

    def populate_serial_numbers(self) -> None:
        if not self.config.serial_numeric_regex:
            return
        self.conn.execute(
            "UPDATE records SET serial_num = TRY_CAST(regexp_extract(serial, ?, 1) AS BIGINT)",
            [self.config.serial_numeric_regex],
        )
